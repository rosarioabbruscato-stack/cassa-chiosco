import os
import sqlite3
import openpyxl
import glob

DB_PATH = "database/negozio.db"

def esporta_magazzino_excel(file_path="export_magazzino.xlsx"):
    """Esporta tutti i prodotti dal database SQLite in un file Excel usando openpyxl."""
    try:
        os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        cursor.execute("SELECT id, codice_barra, nome, prezzo, giacenza FROM prodotti")
        rows = cursor.fetchall()
        column_names = [description[0] for description in cursor.description]
        
        conn.close()
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Magazzino"
        
        ws.append(column_names)
        for row in rows:
            ws.append(row)
            
        wb.save(file_path)
        return True, "Esportazione completata con successo!"
    except Exception as e:
        return False, str(e)

def importa_magazzino_excel(file_path=None):
    """Importa i prodotti da un file Excel nel database SQLite usando openpyxl."""
    if not file_path:
        files_xlsx = glob.glob("*.xlsx")
        if files_xlsx:
            file_path = files_xlsx[0]
        else:
            return False, "Nessun file Excel trovato nella cartella."
            
    try:
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if any(row):
                # Adatta i campi in base alla struttura della tua tabella prodotti
                cursor.execute("""
                    INSERT OR REPLACE INTO prodotti (id, codice_barra, nome, prezzo, giacenza)
                    VALUES (?, ?, ?, ?, ?)
                """, (row[0], row[1], row[2], row[3], row[4]))
                
        conn.commit()
        conn.close()
        return True, "Importazione completata con successo!"
    except Exception as e:
        return False, str(e)
